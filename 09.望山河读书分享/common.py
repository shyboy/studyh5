"""两版读书分享 PPT 共用的版式工具：配色、文字、形状、图片与字号检查。"""
import json
from io import BytesIO
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
IMG = HERE / "images"

MIN_PT = 20

RED = RGBColor(0x9E, 0x1B, 0x1B)
DEEP_RED = RGBColor(0x7A, 0x12, 0x12)
GOLD = RGBColor(0xC8, 0xA1, 0x5A)
LIGHT_GOLD = RGBColor(0xF3, 0xE6, 0xCC)
CREAM = RGBColor(0xFB, 0xF6, 0xEE)
INK = RGBColor(0x2B, 0x2B, 0x2B)
GRAY = RGBColor(0x5A, 0x5A, 0x5A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALE = RGBColor(0xF7, 0xE7, 0xC6)

SANS = "Microsoft YaHei"
KAI = "KaiTi"

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

SLIDES = []  # (title, minutes, notes)
USED_IMAGES = []


def _set_font(run, name):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", name)


def text(slide, x, y, w, h, content, size=24, color=INK, bold=False, font=SANS,
         align="l", anchor="t", em=RED, spacing=1.2, space_after=0):
    """content: 字符串（\\n 分段）或段落列表；用 **…** 标记强调文字。"""
    assert size >= MIN_PT, f"font size {size} < {MIN_PT}"
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = ANCHOR[anchor]
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(0.04))
    paras = content.split("\n") if isinstance(content, str) else content
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN[align]
        p.line_spacing = spacing
        if space_after:
            p.space_after = Pt(space_after)
        for j, seg in enumerate(para.split("**")):
            if not seg:
                continue
            r = p.add_run()
            r.text = seg
            r.font.size = Pt(size)
            r.font.bold = bold or (j % 2 == 1)
            r.font.color.rgb = em if j % 2 == 1 else color
            _set_font(r, font)
        p._p.get_or_add_endParaRPr().set("sz", str(size * 100))
    return box


def add_run(box, content, size, color, bold=False, font=SANS):
    assert size >= MIN_PT
    r = box.text_frame.paragraphs[-1].add_run()
    r.text = content
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    _set_font(r, font)


def shape(slide, kind, x, y, w, h, fill=None, line=None, line_w=1.5):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    s.shadow.inherit = False
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.12
    return s


def rect(slide, x, y, w, h, fill=None, line=None, line_w=1.5):
    return shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, fill, line, line_w)


def rrect(slide, x, y, w, h, fill=None, line=None, line_w=1.5):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, line, line_w)


def oval(slide, x, y, d, fill):
    return shape(slide, MSO_SHAPE.OVAL, x, y, d, d, fill=fill)


def seal(slide, x, y, size, char, fill=RED, color=GOLD, border=GOLD):
    rect(slide, x, y, size, size, fill=fill, line=border, line_w=2.5)
    inset = size * 0.08
    rect(slide, x + inset, y + inset, size - 2 * inset, size - 2 * inset,
         line=border, line_w=1)
    text(slide, x, y, size, size, char, size=max(MIN_PT, int(size * 72 * 0.55)),
         color=color, bold=True, font=KAI, align="c", anchor="m", spacing=1.0)


def arrow(slide, x, y, w=0.9, h=0.6, color=GOLD):
    shape(slide, MSO_SHAPE.RIGHT_ARROW, x, y, w, h, fill=color)


def bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _prepare(name, w, h, focus, shade, dpi=150):
    im = Image.open(IMG / f"{name}.jpg").convert("RGB")
    iw, ih = im.size
    ar = w / h
    if iw / ih > ar:
        nw = int(ih * ar)
        x0 = int((iw - nw) * focus[0])
        im = im.crop((x0, 0, x0 + nw, ih))
    else:
        nh = int(iw / ar)
        y0 = int((ih - nh) * focus[1])
        im = im.crop((0, y0, iw, y0 + nh))
    im = im.resize((int(w * dpi), int(h * dpi)), Image.LANCZOS)
    if shade:
        side, rgb, a0, a1, extent = shade
        W, H = im.size
        horizontal = side in ("left", "right")
        n = W if horizontal else H
        start = int(n * (1 - extent))
        vals = [int(255 * (a0 + (a1 - a0) * (0 if i < start else (i - start) / max(1, n - start))))
                for i in range(n)]
        if side in ("left", "top"):
            vals.reverse()
        strip = Image.new("L", (n, 1) if horizontal else (1, n))
        strip.putdata(vals)
        mask = strip.resize((W, H))
        im = Image.composite(Image.new("RGB", im.size, rgb), im, mask)
    buf = BytesIO()
    im.save(buf, "JPEG", quality=86)
    buf.seek(0)
    return buf


def photo(slide, name, x, y, w, h, focus=(0.5, 0.5), shade=None, frame=None, frame_w=2.5):
    """按框裁切铺满。shade=(方向, RGB, 起始不透明度, 末端不透明度, 渐变占比)，用于在图上压字。"""
    if name not in USED_IMAGES:
        USED_IMAGES.append(name)
    pic = slide.shapes.add_picture(_prepare(name, w, h, focus, shade),
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    if frame is not None:
        pic.line.color.rgb = frame
        pic.line.width = Pt(frame_w)
    return pic


def new_slide(title, minutes, notes):
    slide = prs.slides.add_slide(BLANK)
    SLIDES.append((title, minutes, notes))
    tf = slide.notes_slide.notes_text_frame
    tf.text = f"【建议用时：约 {minutes} 分钟】\n{notes}" if minutes else notes
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.size = Pt(MIN_PT)
            _set_font(r, SANS)
    return slide


def page_no(slide, color=GOLD):
    text(slide, 11.4, 6.85, 1.5, 0.5, f"{len(prs.slides):02d}", size=20, color=color, align="r")


def content_slide(title, label, minutes, notes, x0=0.9, width=11.6):
    s = new_slide(title, minutes, notes)
    bg(s, CREAM)
    rect(s, 0, 0, 0.22, 7.5, fill=RED)
    rect(s, 0.22, 0, 0.06, 7.5, fill=GOLD)
    if label:
        text(s, x0, 0.42, width, 0.45, label, size=20, color=GOLD, bold=True)
    text(s, x0, 0.82, width, 0.85, title, size=36 if width > 9 else 32, color=DEEP_RED, bold=True)
    rect(s, x0, 1.72, 1.2, 0.06, fill=GOLD)
    page_no(s)
    return s


def credits_slide():
    """图片来源页（CC BY-SA 要求署名），放在最后，不计入讲述时间。"""
    credits = json.loads((IMG / "credits.json").read_text(encoding="utf-8"))
    s = new_slide("图片来源", 0, "（本页不讲，仅作图片署名）")
    bg(s, CREAM)
    rect(s, 0, 0, 0.22, 7.5, fill=RED)
    text(s, 0.9, 0.5, 11.5, 0.7, "图片来源", size=32, color=DEEP_RED, bold=True)
    rows = [f"{credits[n]['desc']}：{credits[n]['artist']}，{credits[n]['license']}"
            for n in USED_IMAGES]
    text(s, 0.9, 1.45, 11.5, 5.0, rows, size=20, color=GRAY, spacing=1.05, space_after=4)
    text(s, 0.9, 6.6, 11.5, 0.5, "以上图片均来自 Wikimedia Commons，按原许可协议使用。",
         size=20, color=GRAY)


def check_min_font():
    bad = []
    for idx, slide in enumerate(prs.slides, 1):
        for shp in slide.shapes:
            if not shp.has_text_frame:
                continue
            for p in shp.text_frame.paragraphs:
                for r in p.runs:
                    if r.text.strip() and (r.font.size is None or r.font.size.pt < MIN_PT):
                        bad.append((idx, r.text, r.font.size))
    return bad


def save(pptx_path, script_path, heading, extra=""):
    problems = check_min_font()
    assert not problems, f"字号小于 {MIN_PT}pt 的文字: {problems}"
    prs.save(pptx_path)
    total = sum(m for _, m, _ in SLIDES)
    lines = [f"# {heading}", "", f"共 {len(SLIDES)} 页，建议总用时约 {total:.0f} 分钟{extra}。", ""]
    for i, (title, minutes, notes) in enumerate(SLIDES, 1):
        if not minutes:
            continue
        lines += [f"## 第 {i} 页 · {title}（约 {minutes} 分钟）", "", notes.replace("\n", "\n\n"), ""]
    Path(script_path).write_text("\n".join(lines), encoding="utf-8")
    print(f"saved {Path(pptx_path).name}: {len(SLIDES)} slides, ~{total:.1f} min")
